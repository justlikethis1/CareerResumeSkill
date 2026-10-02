import base64
import json
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree as ET
from pypdf import PdfReader, PdfWriter

from career_resume_skill import docx_cli, docx_engine
from career_resume_skill.docx_engine import (
    DocxLayoutError,
    TextLengthBudgetError,
    _assign_standard_bullet,
    _clean_replacement_text,
    _deduplicate_pdf_pages,
    _sanitize_pdf_dash_fillers,
    _word_conversion_timeout_seconds,
    calculate_docx_layout_budget,
    convert_docx_to_pdf,
    inject_docx_bullet_groups,
    inject_docx_text,
    inspect_docx,
)
from career_resume_skill.docx_profile import (
    _extract_dates,
    _is_role_or_metadata,
    import_docx_as_master_cv,
)
from career_resume_skill.quality import audit_docx_health
from career_resume_skill.tools import find_libreoffice
from career_resume_skill.typography import estimate_string_width

_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _build_fixture(path: Path) -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Candidate Header"
    document.sections[0].footer.paragraphs[0].text = "Candidate Footer"
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Python Engineer")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(31, 78, 121)
    mixed = document.add_paragraph()
    mixed.add_run("Python ")
    mixed_run = mixed.add_run("Engineer")
    mixed_run.bold = True
    leader = document.add_paragraph()
    leader.style = "List Bullet"
    leader.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    leader.paragraph_format.tab_stops.add_tab_stop(
        Inches(5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DASHES
    )
    leader.add_run("Original body").add_tab()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Quant"
    table.cell(0, 1).text = "Research"
    image_path = path.with_suffix(".png")
    image_path.write_bytes(_ONE_PIXEL_PNG)
    document.add_picture(str(image_path), width=Inches(0.2))
    document.save(path)


def test_inspect_docx_extracts_style_grid_and_media(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _build_fixture(source)

    report = inspect_docx(source, tmp_path / "resources")
    paragraph = next(item for item in report.paragraphs if item.text == "Python Engineer")

    assert paragraph.node_id.startswith("p:")
    assert paragraph.runs[0].style["bold"] is True
    assert paragraph.runs[0].style["font_size_pt"] == 11.0
    assert report.layout["columns"] == 1
    assert report.layout["sections"]
    assert "Normal" in report.layout["styles"]
    assert report.layout["styles"]["Normal"]["type"] == "paragraph"
    assert report.layout["stories"]["headers"] == ["Candidate Header"]
    assert report.layout["stories"]["footers"] == ["Candidate Footer"]
    assert any(item.table_path for item in report.paragraphs if item.text == "Quant")
    assert len(report.images) == 1
    assert report.images[0].anchor_type == "inline"
    assert Path(report.images[0].extracted_path).is_file()


def test_inspect_docx_owned_resources_are_cleaned_explicitly(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _build_fixture(source)

    report = inspect_docx(source)
    resources = report.resources_dir
    assert resources.exists()
    report.close()

    assert not resources.exists()


def test_docx_profile_splits_styled_projects_and_generic_institutions(tmp_path: Path) -> None:
    source = tmp_path / "profile.docx"
    document = Document()
    document.add_paragraph("Candidate")
    document.add_paragraph("candidate@example.com")
    document.add_paragraph("Education", style="Heading 1")
    document.add_paragraph("Stanford University")
    document.add_paragraph("MSc Computer Science 2022.09 - 2024.06")
    document.add_paragraph("Tsinghua University")
    document.add_paragraph("BSc Engineering 09.2018 - 06.2022")
    document.add_paragraph("清华大学")
    document.add_paragraph("计算机科学硕士 2022.09 - 2024.06")
    document.add_paragraph("Experience", style="Heading 1")
    document.add_paragraph("Acme Labs Ltd.")
    document.add_paragraph("Software Engineer 01.2023 - 12.2024")
    document.add_paragraph("Projects", style="Heading 1")
    first_title = document.add_paragraph("Latency Profiler", style="Heading 2")
    first_title.add_run(" ").bold = True
    document.add_paragraph("Reduced measured request latency.", style="List Bullet")
    second_title = document.add_paragraph("Dataset Auditor", style="Heading 2")
    second_title.add_run(" ").bold = True
    document.add_paragraph("Validated schema drift in nightly jobs.", style="List Bullet")
    document.add_paragraph("Skills", style="Heading 1")
    document.add_paragraph(
        "Languages: Mandarin (Native), English (Fluent, IELTS 7.0), German (B1)"
    )
    document.save(source)

    imported = import_docx_as_master_cv(source)
    sections = {section["type"]: section for section in imported["master_cv"]["sections"]}

    education = sections["education"]["items"]
    assert [item["organization"] for item in education] == [
        "Stanford University", "Tsinghua University", "清华大学",
    ]
    assert sections["experience"]["items"][0]["organization"] == "Acme Labs Ltd."
    projects = sections["projects"]["items"]
    assert [item["title"] for item in projects] == ["Latency Profiler", "Dataset Auditor"]
    assert projects[0]["bullets"] == ["Reduced measured request latency."]
    assert projects[1]["bullets"] == ["Validated schema drift in nightly jobs."]
    assert imported["master_cv"]["skills"]["Languages"] == [
        "Mandarin (Native)", "English (Fluent, IELTS 7.0)", "German (B1)",
    ]


@pytest.mark.parametrize("date_text", [
    "2021.09 - Present",
    "2021.09 - 至今",
    "Sep 2021 - Jun 2023",
    "2021/09 - 2023/06",
    "2021-9 ~ 2023-6",
    "09.2021 - 06.2023",
])
def test_docx_profile_recognizes_common_date_ranges(date_text: str) -> None:
    assert _is_role_or_metadata(f"Software Engineer {date_text}")
    assert _extract_dates(f"Software Engineer {date_text}") == date_text


@pytest.mark.parametrize("role_text", ["算法工程师", "人工智能研究员", "技术实习生", "数据分析师"])
def test_docx_profile_recognizes_chinese_role_metadata(role_text: str) -> None:
    assert _is_role_or_metadata(role_text)


def test_font_metrics_measure_glyph_width_and_docx_line_budget(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _build_fixture(source)

    budget = calculate_docx_layout_budget(source)

    assert estimate_string_width("WWWW", "Arial", 11) > estimate_string_width("iiii", "Arial", 11)
    assert budget["target_line_occupancy_ratio"] == [0.85, 0.92]
    assert budget["paragraphs"]
    assert budget["paragraphs"][0]["available_width_pt"] > 0
    assert budget["paragraphs"][0]["estimated_source_lines"] >= 1
    assert budget["paragraphs"][0]["bold"] is True
    assert budget["paragraphs"][0]["resolved_font_path"]


def test_cjk_layout_uses_east_asian_run_font(tmp_path: Path) -> None:
    source = tmp_path / "cjk.docx"
    document = Document()
    run = document.add_paragraph().add_run("中文简历经历描述")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    document.save(source)

    budget = calculate_docx_layout_budget(source)

    assert budget["paragraphs"][0]["font_name"] == "Microsoft YaHei"


def test_table_cell_geometry_uses_nearest_cell_width(tmp_path: Path) -> None:
    source = tmp_path / "tables.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).width = Inches(2)
    table.cell(0, 1).width = Inches(4)
    table.cell(0, 0).text = "Narrow column"
    table.cell(0, 1).text = "Wide column"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).width = Inches(1)
    nested.cell(0, 0).text = "Nested cell"
    document.save(source)

    paragraphs = {
        paragraph["source_text"]: paragraph
        for paragraph in calculate_docx_layout_budget(source)["paragraphs"]
    }

    assert paragraphs["Narrow column"]["available_width_pt"] == 144.0
    assert paragraphs["Wide column"]["available_width_pt"] == 288.0
    assert paragraphs["Nested cell"]["available_width_pt"] == 72.0


def test_audit_cv_health_reports_table_and_glyph_risks(tmp_path: Path) -> None:
    source = tmp_path / "audit.docx"
    document = Document()
    document.add_paragraph("Candidate \uf0b7")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python service"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested analysis"
    document.save(source)

    audit = audit_docx_health(source)

    assert audit["estimated_current_lines"] >= 3
    assert audit["estimated_current_line_word_capacity"] > 0
    assert audit["private_use_codepoints"] == ["U+F0B7"]
    assert audit["table_health"]["paragraphs"] >= 2
    assert audit["table_health"]["nested_paragraphs"] >= 1


def test_docx_cli_audit_returns_json_without_jd(tmp_path: Path, monkeypatch, capsys) -> None:
    source = tmp_path / "audit.docx"
    document = Document()
    document.add_paragraph("Candidate \uf0b7")
    document.save(source)
    monkeypatch.setattr("sys.argv", ["career-resume-docx", "audit", str(source)])

    docx_cli.main()

    report = json.loads(capsys.readouterr().out)
    assert report["private_use_codepoints"] == ["U+F0B7"]
    assert report["estimated_current_lines"] >= 1


def test_table_without_width_metadata_uses_column_fallback(tmp_path: Path) -> None:
    source = tmp_path / "unknown-width.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left column"
    table.cell(0, 1).text = "Right column"
    for cell in table.rows[0].cells:
        properties = cell._tc.tcPr
        properties.remove(properties.tcW)
    for column in list(table._tbl.tblGrid):
        table._tbl.tblGrid.remove(column)
    document.save(source)

    budget = calculate_docx_layout_budget(source)
    left = next(entry for entry in budget["paragraphs"] if entry["source_text"] == "Left column")
    audit = audit_docx_health(source)

    assert left["table_cell_width_twips"] is None
    section = document.sections[0]
    assert left["available_width_pt"] == round(
        (section.page_width.pt - section.left_margin.pt - section.right_margin.pt) / 2, 2
    )
    assert audit["table_health"]["widths_unresolved"]


def test_inject_docx_text_preserves_package_and_changes_only_text(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    _build_fixture(source)
    report = inspect_docx(source, tmp_path / "resources")
    paragraph = next(item for item in report.paragraphs if item.text == "Python Engineer")

    inject_docx_text(source, output, {paragraph.node_id: "Quant Engineer"})
    tailored = inspect_docx(output, tmp_path / "tailored-resources")

    changed = next(item for item in tailored.paragraphs if item.node_id == paragraph.node_id)
    assert changed.text == "Quant Engineer."
    assert changed.runs[0].style == paragraph.runs[0].style
    assert len(tailored.images) == 1
    assert tailored.images[0].package_path == report.images[0].package_path


def test_inject_docx_preserves_each_run_style(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    _build_fixture(source)
    report = inspect_docx(source)
    paragraph = next(item for item in report.paragraphs if item.text == "Python Engineer")
    mixed = next(item for item in report.paragraphs if item.text == "Python Engineer" and len(item.runs) == 2)

    inject_docx_text(source, output, {mixed.node_id: "Quant Developer"})
    tailored = inspect_docx(output)
    updated = next(item for item in tailored.paragraphs if item.node_id == mixed.node_id)
    raw_document = Document(output)
    raw_paragraph = next(paragraph for paragraph in raw_document.paragraphs if paragraph.text == "Quant Developer.")

    assert updated.text == "Quant Developer."
    assert len(updated.runs) == 1
    assert len(raw_paragraph.runs) == 2
    assert raw_paragraph.runs[1].text == ""
    assert raw_paragraph.runs[1].bold is True
    assert paragraph.runs[0].style["bold"] is True


def test_inject_docx_rejects_overlong_replacement(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _build_fixture(source)
    report = inspect_docx(source)
    paragraph = next(item for item in report.paragraphs if item.text == "Python Engineer")

    overlong = "A very long technical resume bullet that must exceed the two rendered line budget. " * 30

    with pytest.raises(TextLengthBudgetError, match="rendered lines"):
        inject_docx_text(source, tmp_path / "rejected.docx", {paragraph.node_id: overlong})


def test_inject_docx_removes_tabs_and_leaders_from_replaced_body(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    _build_fixture(source)
    report = inspect_docx(source)
    paragraph = next(item for item in report.paragraphs if item.text == "Original body")

    inject_docx_text(source, output, {paragraph.node_id: "Rewritten body"})
    raw_document = Document(output)
    updated = next(item for item in raw_document.paragraphs if item.text == "Rewritten body.")
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    tab_nodes = list(updated._p.iter(namespace + "tab"))
    assert not tab_nodes
    style = next(updated._p.iter(namespace + "pStyle"))
    alignment = next(updated._p.iter(namespace + "jc"))
    assert style.get(namespace + "val") == "ListBullet"
    assert alignment.get(namespace + "val") == "both"


def test_inject_docx_preserves_custom_indent_when_replacing_bullets(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    document = Document()
    paragraph = document.add_paragraph("Original body")
    paragraph.style = "List Bullet"
    paragraph.paragraph_format.left_indent = Inches(0.75)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    document.save(source)

    report = inspect_docx(source)
    paragraph_record = next(item for item in report.paragraphs if item.text == "Original body")
    inject_docx_text(source, output, {paragraph_record.node_id: "Rewritten body"})

    raw = Document(output)
    updated = next(p for p in raw.paragraphs if p.text == "Rewritten body.")
    assert updated.paragraph_format.left_indent is not None
    assert updated.paragraph_format.left_indent != 0


def test_replacement_uses_body_run_style_after_bold_prefix(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Project: ").bold = True
    paragraph.add_run("Original body")
    document.save(source)
    paragraph_id = next(
        record.node_id for record in inspect_docx(source).paragraphs
        if record.text == "Project: Original body"
    )

    inject_docx_text(source, output, {paragraph_id: "Rewritten result."})
    rewritten = next(p for p in Document(output).paragraphs if p.text == "Rewritten result.")

    assert rewritten.runs[0].bold is not True


def test_standard_bullet_numbering_respects_paragraph_property_order() -> None:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = ET.Element(namespace + "p")
    properties = ET.SubElement(paragraph, namespace + "pPr")
    for name in ("pStyle", "keepNext", "spacing", "ind", "jc", "rPr"):
        ET.SubElement(properties, namespace + name)

    _assign_standard_bullet(paragraph, "7")

    children = [ET.QName(child).localname for child in properties]
    assert children == ["pStyle", "keepNext", "numPr", "spacing", "ind", "jc", "rPr"]


def test_standard_bullet_reorders_existing_numbering_property() -> None:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = ET.Element(namespace + "p")
    properties = ET.SubElement(paragraph, namespace + "pPr")
    for name in ("pStyle", "spacing", "numPr", "ind"):
        node = ET.SubElement(properties, namespace + name)
        if name == "numPr":
            ET.SubElement(node, namespace + "numId").set(namespace + "val", "3")

    _assign_standard_bullet(paragraph, "7")

    children = [ET.QName(child).localname for child in properties]
    assert children == ["pStyle", "numPr", "spacing", "ind"]
    assert properties.find(namespace + "numPr/" + namespace + "numId").get(namespace + "val") == "7"


def test_standard_bullet_inserts_missing_indent_before_alignment_and_run_properties() -> None:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = ET.Element(namespace + "p")
    properties = ET.SubElement(paragraph, namespace + "pPr")
    for name in ("pStyle", "spacing", "jc", "rPr"):
        ET.SubElement(properties, namespace + name)

    _assign_standard_bullet(paragraph, "7")

    children = [ET.QName(child).localname for child in properties]
    assert children == ["pStyle", "numPr", "spacing", "ind", "jc", "rPr"]


def test_inject_docx_rejects_same_input_and_output_path(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _build_fixture(source)

    with pytest.raises(DocxLayoutError, match="must differ"):
        inject_docx_text(source, source, {})


def test_inject_docx_clones_existing_bullet_paragraph_style(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    _build_fixture(source)
    report = inspect_docx(source)
    paragraph = next(item for item in report.paragraphs if item.text == "Original body")

    inject_docx_bullet_groups(
        source,
        output,
        {
            "hero": {
                "paragraph_ids": [paragraph.node_id],
                "bullets": ["First result", "Second result", "Third result"],
            }
        },
    )
    tailored = inspect_docx(output)
    generated = [item.text for item in tailored.paragraphs]

    assert "First result." in generated
    assert "Second result." in generated
    assert "Third result." in generated


def test_company_locations_remain_right_aligned_in_exported_pdf(tmp_path: Path) -> None:
    if find_libreoffice() is None:
        pytest.skip("DOCX PDF alignment check needs LibreOffice")
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    document = Document()
    for company, location, long_spaces in (
        ("Research Lab", "Zhejiang, China", False),
        ("Trading Center", "Hangzhou, China", True),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run(company)
        if long_spaces:
            paragraph.add_run("               ")
        else:
            paragraph.add_run().add_tab()
        paragraph.add_run(" " + location)
    body = document.add_paragraph()
    body.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
    body.add_run("Body with a trailing tab").add_tab()
    document.save(source)

    inject_docx_bullet_groups(source, output, {})
    result = Document(output)
    assert result.paragraphs[0].text == "Research Lab\t Zhejiang, China"
    assert result.paragraphs[1].text == "Trading Center\t Hangzhou, China"
    assert "\t" not in result.paragraphs[2].text
    pdf = convert_docx_to_pdf(output)
    positions: dict[str, float] = {}

    def collect(text, cm, tm, font, size):
        for label in ("Research Lab", "Trading Center", "Zhejiang, China", "Hangzhou, China"):
            if label in text:
                positions[label] = tm[4]

    PdfReader(pdf).pages[0].extract_text(visitor_text=collect)

    assert len(positions) == 4
    assert all(positions[place] - positions[company] > 280 for company, place in (
        ("Research Lab", "Zhejiang, China"),
        ("Trading Center", "Hangzhou, China"),
    ))
    assert abs(positions["Zhejiang, China"] - positions["Hangzhou, China"]) < 20


def test_variable_length_locations_share_right_edge_without_company_collision(tmp_path: Path) -> None:
    if find_libreoffice() is None or shutil.which("pdftotext") is None:
        pytest.skip("PDF word bounding boxes need LibreOffice and pdftotext")
    source = tmp_path / "source.docx"
    document = Document()
    for company, place in (
        ("Research Lab", "China"),
        ("Institute of Flexible Electronics Technology of THU",
         "Jiaxing, Zhejiang, Peoples Republic of China"),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run(company)
        paragraph.add_run().add_tab()
        paragraph.add_run(" " + place)
    document.save(source)
    output = inject_docx_bullet_groups(source, tmp_path / "tailored.docx", {})
    pdf = convert_docx_to_pdf(output)
    result = subprocess.run(
        [shutil.which("pdftotext"), "-bbox", str(pdf), "-"],
        capture_output=True, text=True, check=True,
    )
    words = [word for word in ET.fromstring(result.stdout.encode()).iter()
             if ET.QName(word).localname == "word"]

    def word(label: str):
        return next(value for value in words if value.text == label)

    right_edges = [float(value.get("xMax")) for value in words if value.text == "China"]
    assert len(right_edges) == 2
    assert abs(right_edges[0] - right_edges[1]) < 2
    first_china = next(value for value in words if value.text == "China")
    assert float(word("Research").get("yMin")) == float(first_china.get("yMin"))
    assert float(word("Jiaxing,").get("yMin")) > float(word("THU").get("yMax"))


def test_education_date_stays_beside_degree_below_location(tmp_path: Path) -> None:
    if find_libreoffice() is None:
        pytest.skip("Education PDF alignment check needs LibreOffice")
    source = tmp_path / "education.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("University of Science and Technology" + " " * 18 + "Hong Kong, China" + " " * 90)
    paragraph.add_run("Artificial Intelligence and Entrepreneurship")
    paragraph.add_run().add_tab()
    paragraph.add_run("09.2026-06.2028")
    document.save(source)

    output = inject_docx_bullet_groups(source, tmp_path / "education-output.docx", {})
    education = Document(output).paragraphs[0]
    assert not education._p.xpath(".//w:br")
    pdf = convert_docx_to_pdf(output)
    positions: dict[str, tuple[float, float]] = {}

    def collect(text, cm, tm, font, size):
        for label in ("Artificial Intelligence and Entrepreneurship", "09.2026"):
            if label in text:
                positions[label] = (tm[4], tm[5])

    PdfReader(pdf).pages[0].extract_text(visitor_text=collect)

    assert len(positions) == 2
    degree_x, degree_y = positions["Artificial Intelligence and Entrepreneurship"]
    date_x, date_y = positions["09.2026"]
    assert abs(degree_y - date_y) < 1
    assert date_x - degree_x > 250


def test_empty_bullet_group_removes_unselected_source_bullets(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "tailored.docx"
    document = Document()
    document.add_paragraph("Selected Project", style="Heading 1")
    selected_paragraph = document.add_paragraph("Kept source bullet.", style="List Bullet")
    document.add_paragraph("Omitted Project", style="Heading 1")
    omitted_paragraph = document.add_paragraph("Removed source bullet.", style="List Bullet")
    document.save(source)
    records = inspect_docx(source).paragraphs
    selected_id = next(record.node_id for record in records if record.text == selected_paragraph.text)
    omitted_id = next(record.node_id for record in records if record.text == omitted_paragraph.text)

    inject_docx_bullet_groups(source, output, {
        "selected": {"paragraph_ids": [selected_id], "bullets": ["Kept tailored bullet."]},
        "omitted": {"paragraph_ids": [omitted_id], "bullets": []},
    })

    texts = [record.text for record in inspect_docx(output).paragraphs]
    assert "Kept tailored bullet." in texts
    assert "Removed source bullet." not in texts
    assert "Omitted Project" in texts


def test_empty_bullet_group_preserves_required_paragraph_in_table_cell(tmp_path: Path) -> None:
    source = tmp_path / "table.docx"
    output = tmp_path / "tailored.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "Only source bullet"
    document.save(source)
    paragraph_id = next(
        record.node_id for record in inspect_docx(source).paragraphs
        if record.text == "Only source bullet"
    )

    inject_docx_bullet_groups(source, output, {
        "omitted": {"paragraph_ids": [paragraph_id], "bullets": []},
    })

    result = Document(output)
    result_cell = result.tables[0].cell(0, 0)
    assert len(result_cell.paragraphs) == 1
    assert result_cell.paragraphs[0].text == ""


def test_compact_spacing_includes_heading_and_bullets_but_preserves_fixed_line_spacing(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "compact.docx"
    document = Document()
    heading = document.add_paragraph("Projects")
    heading.paragraph_format.space_after = Pt(2.5)
    bullet = document.add_paragraph("Built a reliable pipeline.")
    bullet.paragraph_format.space_after = Pt(2.5)
    bullet.paragraph_format.line_spacing = 1.0
    fixed = document.add_paragraph("Validated results.")
    fixed.paragraph_format.line_spacing = Pt(12)
    document.save(source)
    paragraph_id = next(
        item.node_id for item in inspect_docx(source).paragraphs
        if item.text == "Built a reliable pipeline."
    )

    inject_docx_bullet_groups(
        source, output,
        {"project": {"paragraph_ids": [paragraph_id], "bullets": ["Built a reliable pipeline."]}},
        compact_spacing=True,
    )
    result = Document(output)

    assert result.paragraphs[0].paragraph_format.space_after.pt < 2.5
    assert result.paragraphs[1].paragraph_format.space_after.pt < 2.5
    assert result.paragraphs[1].paragraph_format.line_spacing < 1.0
    assert result.paragraphs[2].paragraph_format.line_spacing.pt == 12
    assert bullet.paragraph_format.space_after.pt == 2.5


def test_compact_spacing_resolves_inherited_style_and_keeps_heading_with_next(tmp_path: Path) -> None:
    source = tmp_path / "styled.docx"
    output = tmp_path / "compact.docx"
    document = Document()
    document.add_paragraph("Skills", style="Heading 1")
    body_style = document.styles.add_style("CompactBody", 1)
    body_style.paragraph_format.space_after = Pt(6)
    body_style.paragraph_format.line_spacing = 1.0
    body = document.add_paragraph("Languages: Python, C++", style="CompactBody")
    original_text = body.text
    document.save(source)

    inject_docx_bullet_groups(source, output, {}, compact_spacing=True)

    compacted = Document(output)
    heading_properties = compacted.paragraphs[0]._p.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"
    )
    assert heading_properties is not None
    assert heading_properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}keepNext"
    ) is not None
    assert compacted.paragraphs[1].paragraph_format.space_after.pt < 6
    assert compacted.paragraphs[1].paragraph_format.line_spacing < 1.0
    assert compacted.paragraphs[1].text == original_text


def test_heading_gets_minimum_after_spacing_without_replacing_paragraph_boundary(
    tmp_path: Path,
) -> None:
    source = tmp_path / "heading-gap.docx"
    output = tmp_path / "heading-gap-output.docx"
    document = Document()
    document.add_paragraph("PROJECT EXPERIENCE", style="Heading 1")
    document.add_paragraph("Finance_Helper", style="Heading 2")
    document.save(source)

    inject_docx_bullet_groups(source, output, {})

    result = Document(output)
    heading = result.paragraphs[0]
    next_paragraph = result.paragraphs[1]
    assert heading.text == "PROJECT EXPERIENCE"
    assert next_paragraph.text == "Finance_Helper"
    assert heading.paragraph_format.space_after.pt >= 6


def test_uppercase_section_heading_with_generic_style_gets_separation(tmp_path: Path) -> None:
    source = tmp_path / "generic-heading.docx"
    output = tmp_path / "generic-heading-output.docx"
    document = Document()
    document.add_paragraph("PROJECT EXPERIENCE")
    document.add_paragraph("Finance_Helper", style="Heading 2")
    document.save(source)

    inject_docx_bullet_groups(source, output, {})

    result = Document(output)
    assert result.paragraphs[0].paragraph_format.space_after.twips >= 120
    assert result.paragraphs[1].text == "Finance_Helper"


def test_deep_compact_spacing_applies_global_line_pitch_and_margin_floor(tmp_path: Path) -> None:
    source = tmp_path / "styled.docx"
    output = tmp_path / "deep-compact.docx"
    document = Document()
    document.add_paragraph("Heading", style="Heading 1")
    document.add_paragraph("Body paragraph with enough text to exercise deep compaction.")
    document.save(source)

    paragraph_id = next(item.node_id for item in inspect_docx(source).paragraphs if item.text == "Body paragraph with enough text to exercise deep compaction.")
    inject_docx_bullet_groups(
        source,
        output,
        {"body": {"paragraph_ids": [paragraph_id], "bullets": ["Body paragraph with enough text to exercise deep compaction."]}},
        compact_level=2,
    )

    compacted = Document(output)
    assert compacted.paragraphs[1].paragraph_format.line_spacing < 1.0
    section = compacted.sections[0]
    assert section.top_margin.twips >= 540
    assert section.bottom_margin.twips >= 540


def test_deep_compact_spacing_preserves_heading_breathing_room(tmp_path: Path) -> None:
    source = tmp_path / "heading.docx"
    output = tmp_path / "heading-deep.docx"
    document = Document()
    heading = document.add_paragraph("EXPERIENCE", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(0)
    document.add_paragraph("A compact body paragraph.")
    document.save(source)

    inject_docx_bullet_groups(source, output, {}, compact_level=2)

    compacted = Document(output)
    heading = compacted.paragraphs[0].paragraph_format
    assert heading.space_before.twips >= 60
    assert heading.space_after.twips >= 30


def test_convert_docx_to_pdf_with_libreoffice(tmp_path: Path) -> None:
    if find_libreoffice() is None:
        pytest.skip("LibreOffice is not installed")
    source = tmp_path / "source.docx"
    _build_fixture(source)

    pdf = convert_docx_to_pdf(source, tmp_path / "pdf")

    assert pdf.exists()


def test_libreoffice_conversion_uses_headless_flags_and_isolated_profile(
    monkeypatch, tmp_path: Path
) -> None:
    source = tmp_path / "resume.docx"
    source.write_bytes(b"docx")
    captured: dict[str, object] = {}

    def fake_run(command, **kwargs):
        captured["command"] = command
        captured.update(kwargs)
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / "resume.pdf").write_bytes(b"pdf")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(docx_engine, "find_microsoft_word", lambda: None)
    monkeypatch.setattr(docx_engine, "find_libreoffice", lambda: "soffice")
    monkeypatch.setattr(docx_engine.subprocess, "run", fake_run)

    result = convert_docx_to_pdf(source, tmp_path / "pdf")

    command = captured["command"]
    assert result == tmp_path / "pdf" / "resume.pdf"
    assert all(flag in command for flag in (
        "--headless", "--invisible", "--nologo", "--nodefault",
        "--nofirststartwizard", "--nolockcheck",
    ))
    profile_arg = next(arg for arg in command if arg.startswith("-env:UserInstallation="))
    profile_uri = profile_arg.split("=", 1)[1]
    assert profile_uri.startswith("file:")
    assert captured["timeout"] == 180
    assert result.stat().st_size > 0


def test_word_timeout_falls_back_to_libreoffice_and_replaces_partial_pdf(
    monkeypatch, tmp_path: Path
) -> None:
    source = tmp_path / "resume.docx"
    source.write_bytes(b"docx")
    output_dir = tmp_path / "pdf"
    word_calls: list[str] = []
    libreoffice_calls: list[str] = []

    def fake_word(source_path: Path, pdf_path: Path) -> Path:
        word_calls.append(str(source_path))
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        pdf_path.write_bytes(b"partial Word PDF")
        raise subprocess.TimeoutExpired("winword", 60)

    def fake_libreoffice(source_path: Path, target_dir: Path, pdf_path: Path, soffice: str) -> Path:
        libreoffice_calls.append(soffice)
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        target_dir.mkdir(parents=True, exist_ok=True)
        with pdf_path.open("wb") as output:
            writer.write(output)
        return pdf_path

    monkeypatch.setattr(docx_engine, "find_microsoft_word", lambda: "winword")
    monkeypatch.setattr(docx_engine, "find_libreoffice", lambda: "soffice")
    monkeypatch.setattr(docx_engine, "_convert_docx_with_word", fake_word)
    monkeypatch.setattr(docx_engine, "_convert_docx_with_libreoffice", fake_libreoffice)

    result = convert_docx_to_pdf(source, output_dir)

    assert result == output_dir / "resume.pdf"
    assert len(PdfReader(result).pages) == 1
    assert word_calls == [str(source.resolve())]
    assert libreoffice_calls == ["soffice"]


@pytest.mark.parametrize(("configured", "expected"), [
    (None, 60.0),
    ("30", 30.0),
    ("0", 1.0),
    ("500", 180.0),
    ("invalid", 60.0),
])
def test_word_conversion_timeout_is_bounded(monkeypatch, configured: str | None, expected: float) -> None:
    if configured is None:
        monkeypatch.delenv("CAREER_SKILL_WORD_CONVERSION_TIMEOUT_SECONDS", raising=False)
    else:
        monkeypatch.setenv("CAREER_SKILL_WORD_CONVERSION_TIMEOUT_SECONDS", configured)

    assert _word_conversion_timeout_seconds() == expected


def test_clean_replacement_text_removes_padding_not_valid_hyphens() -> None:
    assert _clean_replacement_text("• Built an out-of-domain 16-bit pipeline.-----") == (
        "Built an out-of-domain 16-bit pipeline."
    )
    assert _clean_replacement_text("使用 Python 优化吞吐量 10000 QPS.") == (
        "使用 Python 优化吞吐量 10000 QPS。"
    )
    assert _clean_replacement_text("Built a verified pipeline.---") == (
        "Built a verified pipeline."
    )
    assert _clean_replacement_text("Built a pipeline.\u200b_\t") == "Built a pipeline."


def test_sanitize_pdf_dash_fillers_removes_only_standalone_dash_objects(tmp_path: Path) -> None:
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject

    pdf = tmp_path / "dash.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    stream = DecodedStreamObject()
    stream.set_data(b"BT [(normal-text)] TJ [(-----)] TJ [(-)7(-)7(-)7(-)-3(-)] TJ ET")
    page.replace_contents(stream)
    with pdf.open("wb") as stream:
        writer.write(stream)
    original_page_count = len(PdfReader(str(pdf)).pages)

    _sanitize_pdf_dash_fillers(pdf)

    assert b"-----" not in pdf.read_bytes()
    assert b"normal-text" in pdf.read_bytes()
    assert len(PdfReader(str(pdf)).pages) == original_page_count


def test_deduplicate_pdf_pages_removes_adjacent_exact_text_duplicate(tmp_path: Path) -> None:
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject

    pdf = tmp_path / "duplicate.pdf"
    writer = PdfWriter()
    for _ in range(2):
        page = writer.add_blank_page(width=100, height=100)
        stream = DecodedStreamObject()
        stream.set_data(b"BT [(same page)] TJ ET")
        page.replace_contents(stream)
    with pdf.open("wb") as stream:
        writer.write(stream)

    _deduplicate_pdf_pages(pdf)

    assert len(PdfReader(str(pdf)).pages) == 1
