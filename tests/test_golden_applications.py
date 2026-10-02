import asyncio
import json
import re
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from career_resume_skill.application import ApplicationService
from career_resume_skill.config import Settings
from career_resume_skill.quality import extract_pdf_text
from career_resume_skill.tools import find_libreoffice, find_tectonic


class _Offline:
    available = False
    model_name = "deterministic_fallback"


@pytest.mark.parametrize("role,jd", [
    ("quant_cplusplus_hft", "Quant C++ low latency Python market systems"),
    ("ai_lab_rlhf_dpo", "AI Lab RLHF DPO PyTorch alignment research"),
    ("systems_distributed_backend", "Distributed Python systems high throughput backend"),
])
def test_golden_docx_application_package(tmp_path: Path, role: str, jd: str) -> None:
    if find_libreoffice() is None or find_tectonic() is None:
        pytest.skip("Golden PDF test needs LibreOffice and Tectonic")
    source = tmp_path / "source.docx"
    document = Document()
    for line in (
        "Candidate", "candidate@example.com", "PROJECTS", "Verified Research",
        "Published IEEE ICEICT work using Python with a 5.2% gain and 12.5% error.",
        "Built a BERT pipeline with sub-10ms intent recognition.",
    ):
        document.add_paragraph(line)
    document.save(source)
    service = ApplicationService(Settings(output_dir=str(tmp_path)), provider=_Offline())

    package = asyncio.run(service.generate_docx_application_package(
        str(source), jd, str(tmp_path / f"{role}.docx"), f"Example {role}",
        application_date="2026-04-08",
    ))
    _, pdf_text = extract_pdf_text(package["output_pdf"])
    _, cover_pdf_text = extract_pdf_text(package["cover_letter_pdf"]["pdf_path"])

    assert len(PdfReader(package["output_pdf"]).pages) == 1
    assert len(PdfReader(package["cover_letter_pdf"]["pdf_path"]).pages) == 1
    assert all(term in pdf_text for term in ("5.2%", "12.5%", "sub-10ms", "IEEE ICEICT"))
    assert not any("\ue000" <= char <= "\uf8ff" for char in pdf_text)
    assert not re.search(r"(?:--|[—–])\s*$", pdf_text, flags=re.MULTILINE)
    assert len(package["cover_letter"]["paragraphs"]) == 3
    cover_text = Path(package["cover_letter_text_path"]).read_text(encoding="utf-8")
    assert "Date: April 8, 2026" in cover_text
    assert "Hiring Committee\n" + f"Example {role}" in cover_text
    assert "RE: Application for " in cover_text
    assert "Sincerely,\n\n\n\nCandidate" in cover_text
    assert "Date: April 8, 2026" in Path(package["cover_letter_md_path"]).read_text(encoding="utf-8")
    assert "April 8, 2026" in cover_pdf_text
    assert package["preflight"]["passed"]
    assert package["refinement"]["layout_heuristic"]["initial_compact_level"] == 0
    assert package["refinement"]["applied_compact_level"] == 0
    assert package["performance"]["total_seconds"] >= 0
    assert "resume_injection_and_pdf_attempts" in package["performance"]["stages_seconds"]
    assert package["quality_report"]["integrity_linter"]["passed"]
    assert package["verified"] == (
        package["quality_report"]["integrity_linter"]["passed"]
        and all(
            package["quality_report"][name]["passed"]
            for name in (
                "cover_letter_critic", "cover_letter_content_critic",
                "resume_content_critic", "docx_critic",
            )
        )
    )
    if not package["verified"]:
        assert package["ledger_entry"] is None
    typography = package["quality_report"]["typography_audit"]
    assert typography["final_page_count"] == 1
    assert typography["cover_letter_word_count"] > 0
    assert typography["cl_within_target_range"] is False
    assert Path(package["output_docx"]).name == f"{role}.docx"
    assert Path(package["output_pdf"]).name == f"{role}.pdf"
    assert not package["quality_report"]["ats_readability"]["ats_readability_flags"]


@pytest.mark.parametrize("jd,experience,expected,unsupported", [
    (
        "Finance operations analyst: accounts payable reconciliation, Excel, SAP reporting",
        (
            "Reconciled 120 invoices in Excel and reduced month-end exceptions by 8%.",
            "Prepared monthly financial reporting with documented audit checks.",
        ),
        ("120", "8%", "Excel"), "SAP",
    ),
    (
        "Customer support specialist: resolve complaints in Salesforce CRM",
        (
            "Resolved 95 customer cases with documented triage and follow-up.",
            "Reduced response time by 20% using Zendesk case workflows.",
        ),
        ("95", "20%", "Zendesk"), "Salesforce",
    ),
])
def test_golden_docx_other_candidate_backgrounds(
    tmp_path: Path, jd: str, experience: tuple[str, ...],
    expected: tuple[str, ...], unsupported: str,
) -> None:
    if find_libreoffice() is None or find_tectonic() is None:
        pytest.skip("Golden PDF test needs LibreOffice and Tectonic")
    source = tmp_path / "source.docx"
    document = Document()
    for line in ("Candidate", "candidate@example.com", "EXPERIENCE", "Verified Work", *experience):
        document.add_paragraph(line)
    document.save(source)
    service = ApplicationService(Settings(output_dir=str(tmp_path)), provider=_Offline())

    package = asyncio.run(service.generate_docx_application_package(
        str(source), jd, str(tmp_path / "tailored.docx"), "Example Company",
    ))
    resume_pages, resume_text = extract_pdf_text(package["output_pdf"])
    cover_pages, _ = extract_pdf_text(package["cover_letter_pdf"]["pdf_path"])

    assert (resume_pages, cover_pages) == (1, 1)
    assert all(term in resume_text for term in expected)
    assert unsupported not in resume_text
    assert package["refinement"]["layout_heuristic"]["initial_compact_level"] == 0
    assert package["quality_report"]["integrity_linter"]["passed"]
    assert len(package["cover_letter"]["paragraphs"]) == 3
    assert package["ledger_entry"] is None or package["verified"]


def test_golden_finance_latex_application(tmp_path: Path) -> None:
    if find_tectonic() is None:
        pytest.skip("Golden PDF test needs Tectonic")
    source = Path(__file__).resolve().parents[1] / "examples" / "finance_operations_cv.synthetic.json"
    master_cv = json.loads(source.read_text(encoding="utf-8"))
    service = ApplicationService(Settings(output_dir=str(tmp_path)), provider=_Offline())

    package = asyncio.run(service.generate_application_package(
        "Finance operations analyst: invoice reconciliation, Excel, audit-ready "
        "month-end reporting; SAP preferred.", master_cv, "Example Company",
    ))
    resume_pages, resume_text = extract_pdf_text(package["documents"]["resume"]["pdf_path"])
    cover_pages, _ = extract_pdf_text(package["documents"]["cover_letter"]["pdf_path"])

    assert (resume_pages, cover_pages) == (1, 1)
    assert all(term in resume_text for term in ("120", "8%", "Excel"))
    assert not any(line.endswith("-") for line in resume_text.splitlines())
    assert "Overfull" not in package["documents"]["resume"]["log"]
    assert "SAP" not in resume_text
    assert package["jd_analysis"]["role_type"] == "finance_operations"
    assert package["ats_report"]["integrity_linter"]["resume"]["passed"]
    assert len(package["cover_letter"]["paragraphs"]) == 3
    assert package["ats_report"]["ledger_entry"] is None or package["verified"]


def test_finance_live_shaped_model_package_verifies_both_pdfs(tmp_path: Path) -> None:
    if find_tectonic() is None:
        pytest.skip("Full model-shaped PDF check needs Tectonic")

    class _FinanceModel:
        available = True
        model_name = "synthetic_test_model"

        async def complete_json(self, task, payload, response_model, temperature=0.2):
            if response_model.__name__ == "JDAnalysis":
                return response_model(role_type="finance_operations", must_haves=["Excel"])
            if response_model.__name__ == "GeneratedTailoring":
                return response_model.model_validate({
                    "profile": "Finance operations analyst with documented invoice reconciliation and month-end reporting experience.",
                    "profile_evidence_ids": ["profile"],
                    "sections": [{"type": "experience", "title": "Experience", "items": [{
                        "source_item_id": "section:0:item:0", "bullets": [
                            {"text": "Reconciled 120 supplier invoices each month in Excel against purchase orders and payment records.",
                             "evidence_ids": ["section:0:item:0:bullet:0"]},
                            {"text": "Reduced month-end reconciliation exceptions by 8% through documented checks.",
                             "evidence_ids": ["section:0:item:0:bullet:1"]},
                            {"text": "Prepared financial reports with records for internal audit review.",
                             "evidence_ids": ["section:0:item:0:bullet:2"]},
                        ],
                    }]}], "skills": {"Tools": ["Excel"]}, "evidence_gaps": [], "strategy": {},
                })
            sentences = (
                "My invoice reconciliation in Excel supports traceable finance operations.",
                "Documented validation checks connect payment reviews with recorded monthly reporting.",
                "Careful audit preparation helps teams review supported financial information.",
            )
            return response_model.model_validate({
                "salutation": "Dear Hiring Team,",
                "paragraphs": [" ".join([sentence] * count) for sentence, count in zip(sentences, (6, 16, 10))],
                "closing": "Sincerely", "evidence_ids": ["section:0:item:0:bullet:0"],
            })

    source = Path(__file__).resolve().parents[1] / "examples" / "finance_operations_cv.synthetic.json"
    master_cv = json.loads(source.read_text(encoding="utf-8"))
    service = ApplicationService(Settings(output_dir=str(tmp_path)), provider=_FinanceModel())
    package = asyncio.run(service.generate_application_package(
        "Finance operations analyst: reconcile supplier invoices in Excel and prepare audit-ready reports; SAP preferred.",
        master_cv, "Example Company",
    ))
    resume_pages, resume_text = extract_pdf_text(package["documents"]["resume"]["pdf_path"])
    cover_pages, cover_text = extract_pdf_text(package["documents"]["cover_letter"]["pdf_path"])

    assert package["verified"] and package["ats_report"]["ledger_entry"]
    assert (resume_pages, cover_pages) == (1, 1)
    assert "120" in resume_text and "8%" in resume_text
    assert "SAP" not in resume_text and "Excel" in cover_text


def test_finance_latex_restores_model_only_duration_before_pdf(tmp_path: Path) -> None:
    if find_tectonic() is None:
        pytest.skip("Golden PDF test needs Tectonic")

    class _PersistentDuration:
        available = True
        model_name = "synthetic_test_model"

        async def complete_json(self, task, payload, response_model, temperature=0.2):
            if response_model.__name__ == "JDAnalysis":
                return response_model(role_type="finance_operations", must_haves=["Excel"])
            return response_model.model_validate({
                "profile": "Finance operations analyst with 2 years of invoice reconciliation.",
                "profile_evidence_ids": ["profile"],
                "sections": [{"type": "experience", "title": "Experience", "items": [{
                    "source_item_id": "section:0:item:0",
                    "bullets": [{
                        "text": "Reconciled 120 supplier invoices in Excel over 2 years.",
                        "evidence_ids": ["section:0:item:0:bullet:0"],
                    }],
                }]}], "skills": {"Tools": ["Excel"]}, "evidence_gaps": [], "strategy": {},
            })

    source = Path(__file__).resolve().parents[1] / "examples" / "finance_operations_cv.synthetic.json"
    master_cv = json.loads(source.read_text(encoding="utf-8"))
    service = ApplicationService(Settings(output_dir=str(tmp_path)), provider=_PersistentDuration())

    async def offline_cover(cv, company, tone, context, jd):
        from career_resume_skill.tailoring import draft_cover
        return await draft_cover(cv, company, tone, _Offline(), context, jd)

    service.draft_cover_letter = offline_cover
    package = asyncio.run(service.generate_application_package(
        "Finance operations analyst: Excel invoice reconciliation; SAP preferred.",
        master_cv, "Example Company",
    ))
    resume_pages, resume_text = extract_pdf_text(package["documents"]["resume"]["pdf_path"])

    assert resume_pages == 1
    assert "120" in resume_text and "8%" in resume_text
    assert "2 years" not in resume_text and "SAP" not in resume_text
    assert package["ats_report"]["integrity_linter"]["resume"]["passed"]
    assert any("Restored source profile" in gap for gap in package["tailored_cv"]["evidence_gaps"])
    assert any("Restored source bullet" in gap for gap in package["tailored_cv"]["evidence_gaps"])


def test_golden_customer_support_latex_application(tmp_path: Path) -> None:
    if find_tectonic() is None:
        pytest.skip("Golden PDF test needs Tectonic")
    master_cv = {
        "name": "Example Candidate", "contact": {"email": "example@example.com"},
        "profile": "Customer support specialist focused on documented case handling.",
        "sections": [{"type": "experience", "title": "Experience", "items": [{
            "title": "Customer Support Specialist", "organization": "Example Service",
            "dates": "2023--2026", "bullets": [
                "Resolved 95 customer cases with documented triage and follow-up.",
                "Reduced response time by 20% using Zendesk case workflows.",
            ],
        }]}], "skills": {"Tools": ["Zendesk"]},
    }
    service = ApplicationService(Settings(output_dir=str(tmp_path)), provider=_Offline())

    package = asyncio.run(service.generate_application_package(
        "Customer support specialist: resolve complaints and track cases in Salesforce CRM; "
        "respond promptly and document follow-up.", master_cv, "Example Company",
    ))
    resume_pages, resume_text = extract_pdf_text(package["documents"]["resume"]["pdf_path"])
    cover_pages, _ = extract_pdf_text(package["documents"]["cover_letter"]["pdf_path"])

    assert (resume_pages, cover_pages) == (1, 1)
    assert all(term in resume_text for term in ("95", "20%", "Zendesk"))
    assert "Salesforce" not in resume_text
    assert package["jd_analysis"]["role_type"] == "customer_service"
    assert package["ats_report"]["integrity_linter"]["resume"]["passed"]
    assert len(package["cover_letter"]["paragraphs"]) == 3
    assert package["ats_report"]["ledger_entry"] is None or package["verified"]


def test_chinese_resume_docx_pdf_preserves_evidence(tmp_path: Path) -> None:
    if find_libreoffice() is None or find_tectonic() is None:
        pytest.skip("Chinese PDF test needs LibreOffice and Tectonic")

    class _ChineseModel:
        available = True
        model_name = "synthetic_test_model"

        async def complete_json(self, task, payload, response_model, temperature=0.2):
            if response_model.__name__ == "JDAnalysis":
                return response_model(role_type="ai_research", must_haves=["Python"])
            if response_model.__name__ == "GeneratedTailoring":
                return response_model.model_validate({
                    "profile": "", "profile_evidence_ids": [],
                    "sections": [{"type": "projects", "title": "Projects", "items": [{
                        "source_item_id": "p:3:body",
                        "bullets": [{"text": "使用 Python 完成模型评估，将错误率降低 5.2%。",
                                     "evidence_ids": ["p:3:body:bullet:0"]}],
                    }]}], "skills": {}, "evidence_gaps": [], "strategy": {},
                })
            return response_model.model_validate({
                "salutation": "Dear Hiring Team,",
                "paragraphs": [
                    "I am interested in this technical role and bring verified Python experience relevant to its stated engineering needs. " * 3,
                    "My documented work includes evaluating models, improving error by 5.2%, preserving reproducibility, and building reliable technical workflows. " * 9,
                    "I value careful engineering and collaborative problem solving. I welcome a technical discussion about the role and its priorities. " * 4,
                ],
                "closing": "Sincerely, Candidate",
                "evidence_ids": ["p:3:body:bullet:0"], "evidence_gaps": [],
            })

    source = tmp_path / "source.docx"
    document = Document()
    for line in ("Candidate", "candidate@example.com", "PROJECTS", "Model Evaluation",
                 "Used Python to evaluate models and reduced error by 5.2%."):
        document.add_paragraph(line)
    document.save(source)
    service = ApplicationService(Settings(output_dir=str(tmp_path)), provider=_ChineseModel())

    package = asyncio.run(service.generate_docx_application_package(
        str(source), "AI Lab Python evaluation", str(tmp_path / "chinese.docx"),
        "Example", resume_language="zh_CN",
    ))
    _, text = extract_pdf_text(package["output_pdf"])

    assert "使用 Python" in text
    assert "5.2%" in text
    assert package["resume_language"] == "zh_CN"
    assert package["quality_report"]["pages"]["output"] == 1